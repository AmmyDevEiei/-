import docx
from docx.shared import Pt, Inches, RGBColor
from docx.oxml.ns import qn, nsdecls
from docx.oxml import OxmlElement, parse_xml
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
import os
import sys

sys.stdout.reconfigure(encoding='utf-8')

# Open document
doc = docx.Document(r'แผนที่สมบูรณ์สำหรับแทนนี่/025 ณัฐณิชา งานเล่มคู่มือผลิตสื่อ Video.docx')
print('Initial paragraphs:', len(doc.paragraphs))

# Keep P0 to P37, remove P38 onwards
body = doc._body._element
for p in list(doc.paragraphs[38:]):
    body.remove(p._element)

print('Paragraphs after removal:', len(doc.paragraphs))
doc.save('scratch/test_truncated.docx')
print('Saved test truncated successfully!')
